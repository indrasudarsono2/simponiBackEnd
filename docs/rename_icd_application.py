from pathlib import Path
from docx import Document

path = Path(__file__).parent / "SIMPONI-e-chain-ICD-Draft.docx"
doc = Document(path)


def replace_in_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = run.text.replace("SIMPONI", "PERFORMA").replace("Simponi", "Performa").replace("simponi", "performa")


def replace_in_table(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(paragraph)
            for nested in cell.tables:
                replace_in_table(nested)


for paragraph in doc.paragraphs:
    replace_in_paragraph(paragraph)
for table in doc.tables:
    replace_in_table(table)
for section in doc.sections:
    for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
        for paragraph in part.paragraphs:
            replace_in_paragraph(paragraph)
        for table in part.tables:
            replace_in_table(table)

props = doc.core_properties
props.title = props.title.replace("SIMPONI", "PERFORMA")
props.subject = props.subject.replace("SIMPONI", "PERFORMA")
props.author = props.author.replace("SIMPONI", "PERFORMA")
props.comments = props.comments.replace("SIMPONI", "PERFORMA")

new_path = path.with_name("PERFORMA-e-chain-ICD-Draft.docx")
doc.save(new_path)
print(new_path)
